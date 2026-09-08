import openpyxl
import json
import uuid
import datetime
import os

excel_file = 'normalizacion_datos.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)
ws = wb['Hoja1']

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

# Mapeo oficial de los 4 sectores
sectores_map = {
    'paso lateral': {
        'id': '11111111-0000-0000-0000-000000000001',
        'codigo_sector': 'SEC-PASO',
        'nombre_sector': 'Sector Paso Lateral',
        'descripcion': 'Zona paso lateral y vías perimetrales'
    },
    'san jose': {
        'id': '11111111-0000-0000-0000-000000000002',
        'codigo_sector': 'SEC-SANJOSE',
        'nombre_sector': 'Sector San José',
        'descripcion': 'Barrio San José y ramal alto'
    },
    'centro parroquial': {
        'id': '11111111-0000-0000-0000-000000000003',
        'codigo_sector': 'SEC-CENTRO',
        'nombre_sector': 'Sector Centro Parroquial',
        'descripcion': 'Casco central y parque principal'
    },
    'san antonio': {
        'id': '11111111-0000-0000-0000-000000000004',
        'codigo_sector': 'SEC-SANANTONIO',
        'nombre_sector': 'Sector San Antonio',
        'descripcion': 'Barrio San Antonio y zona baja'
    }
}

NAMESPACE_SIGA = uuid.UUID('6ba7b810-9dad-11d1-80b4-00c04fd430c8')
ADMIN_ID = '00000000-0000-0000-0000-000000000001'

PERIODO_JULIO_ID = '33333333-0000-0000-0000-000000000000'
PERIODO_AGOSTO_ID = '33333333-0000-0000-0000-000000000001'

periodos = [
    {
        'id': PERIODO_JULIO_ID,
        'periodo_codigo': '2026-07',
        'nombre': 'Período Julio 2026',
        'fecha_inicio': '2026-07-01',
        'fecha_fin': '2026-07-31',
        'estado': 'CERRADO'
    },
    {
        'id': PERIODO_AGOSTO_ID,
        'periodo_codigo': '2026-08',
        'nombre': 'Período Agosto 2026',
        'fecha_inicio': '2026-08-01',
        'fecha_fin': '2026-08-31',
        'estado': 'ABIERTO'
    }
]

socios_dict = {}
medidores_list = []
lecturas_list = []
facturas_list = []
multas_list = []

socios_counter = 1
sn_counter = 1

rows = []
for r in range(2, ws.max_row + 1):
    row_data = {headers[c-1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
    if any(v is not None for v in row_data.values()):
        row_data['_row_idx'] = r
        rows.append(row_data)

print(f"Filas leídas de Hoja1: {len(rows)}")

for r in rows:
    row_idx = r['_row_idx']
    nombres = str(r['nombres']).strip() if r['nombres'] is not None else ''
    apellidos = str(r['apellidos']).strip() if r['apellidos'] is not None else ''
    
    # Manejo de cédula: caso especial Patricia Flores (fila 43) socia independiente
    if row_idx == 43 and 'Patricia' in nombres:
        cedula = '1800000081'
    else:
        raw_ced = str(r['cedula_ruc']).strip()
        if '.' in raw_ced:
            raw_ced = raw_ced.split('.')[0]
        cedula = raw_ced.zfill(10)
        
    fnac_val = r['fecha_nacimiento']
    if isinstance(fnac_val, (datetime.datetime, datetime.date)):
        fecha_nacimiento = fnac_val.strftime('%Y-%m-%d')
    else:
        fecha_nacimiento = '1970-01-01'
        
    funion_val = r['fecha_union']
    if isinstance(funion_val, (datetime.datetime, datetime.date)):
        fecha_union = funion_val.strftime('%Y-%m-%d')
    else:
        fecha_union = '2022-01-01'
        
    tel_raw = str(r['telefono']).strip() if r['telefono'] is not None else ''
    if tel_raw and tel_raw.lower() != 'none':
        telefono = tel_raw.split('.')[0]
    else:
        telefono = '0990000000'
        
    sector_key = str(r['sector']).strip().lower()
    sector_info = sectores_map.get(sector_key, sectores_map['centro parroquial'])
    
    dir_raw = str(r['direccion']).strip() if r['direccion'] is not None else ''
    if dir_raw and dir_raw.lower() != 'none':
        direccion = dir_raw
    else:
        direccion = f"Sector {sector_info['nombre_sector']}"
        
    # Registrar Socio si no existe
    if cedula not in socios_dict:
        socio_id = str(uuid.uuid5(NAMESPACE_SIGA, f"socio_{cedula}"))
        codigo_socio = f"SOC-{str(socios_counter).zfill(4)}"
        socios_counter += 1
        
        # Calcular edad
        try:
            nac_year = int(fecha_nacimiento.split('-')[0])
            es_3ra_edad = (2026 - nac_year) >= 65
        except:
            es_3ra_edad = False
            
        socios_dict[cedula] = {
            'id': socio_id,
            'codigo_socio': codigo_socio,
            'nombres': nombres,
            'apellidos': apellidos,
            'cedula_ruc': cedula,
            'fecha_nacimiento': fecha_nacimiento,
            'fecha_union': fecha_union,
            'telefono': telefono,
            'direccion': direccion,
            'id_sector': sector_info['id'],
            'sector_nombre': sector_info['nombre_sector'],
            'estado': 'ACTIVO',
            'es_tercera_edad': es_3ra_edad,
            'medidores_count': 0,
            'deuda_total_acumulada': 0.0,
            'multas_total_acumulada': 0.0
        }
        
    socio = socios_dict[cedula]
    socio['medidores_count'] += 1
    
    # Manejo de número de medidor
    raw_med = str(r['numero_medidor']).strip() if r['numero_medidor'] is not None else ''
    if raw_med.upper() == 'SN' or not raw_med or raw_med.lower() == 'none':
        numero_medidor = f"SN-{str(sn_counter).zfill(2)}"
        sn_counter += 1
    else:
        if '.' in raw_med and raw_med.replace('.', '').isdigit():
            numero_medidor = raw_med.split('.')[0]
        else:
            numero_medidor = raw_med
            
    # Alias de medidor
    alias_raw = str(r['alias_medidor']).strip() if r['alias_medidor'] is not None else ''
    if alias_raw and alias_raw.lower() != 'none':
        alias = alias_raw
    else:
        alias = 'Casa principal' if socio['medidores_count'] == 1 else f"Medidor {socio['medidores_count']}"
        
    tiene_alc = str(r['tiene_alcantarillado']).strip().upper() == 'SI'
    
    try:
        lectura_ini = float(r['lectura_inicial']) if r['lectura_inicial'] is not None else 0.0
    except:
        lectura_ini = 0.0
        
    try:
        deuda_pen = round(float(r['deuda_pendiente']), 2) if r['deuda_pendiente'] is not None else 0.0
    except:
        deuda_pen = 0.0
        
    try:
        multas_val = round(float(r['multas']), 2) if r['multas'] is not None else 0.0
    except:
        multas_val = 0.0
        
    medidor_id = str(uuid.uuid5(NAMESPACE_SIGA, f"medidor_{numero_medidor}"))
    
    medidor_data = {
        'id': medidor_id,
        'id_socio': socio['id'],
        'id_sector': sector_info['id'],
        'numero_medidor': numero_medidor,
        'alias': alias,
        'direccion': direccion,
        'tiene_alcantarillado': tiene_alc,
        'estado': 'ACTIVO',
        'lectura_inicial': lectura_ini,
        'deuda_pendiente': deuda_pen,
        'multas': multas_val,
        'socio_codigo': socio['codigo_socio'],
        'socio_nombre': f"{socio['nombres']} {socio['apellidos']}",
        'sector_nombre': sector_info['nombre_sector']
    }
    medidores_list.append(medidor_data)
    
    socio['deuda_total_acumulada'] = round(socio['deuda_total_acumulada'] + deuda_pen, 2)
    socio['multas_total_acumulada'] = round(socio['multas_total_acumulada'] + multas_val, 2)
    
    # 1. Lectura Base de Corte (Julio 2026)
    # Permite que en Agosto la lectura anterior sea exactamente la lectura inicial
    lectura_id = str(uuid.uuid5(NAMESPACE_SIGA, f"lectura_julio_{numero_medidor}"))
    lecturas_list.append({
        'id': lectura_id,
        'id_medidor': medidor_id,
        'id_socio': socio['id'],
        'id_periodo': PERIODO_JULIO_ID,
        'lectura_anterior': lectura_ini,
        'lectura_actual': lectura_ini,
        'consumo_total': 0.0,
        'excedente_m3': 0.0,
        'fecha_lectura': '2026-07-31T23:59:59.000Z',
        'id_lector': ADMIN_ID,
        'observaciones': 'Lectura de corte inicial - Julio 2026',
        'version': 1
    })
    
    # 2. Factura de Deuda Pendiente al corte de Julio
    if deuda_pen > 0:
        factura_id = str(uuid.uuid5(NAMESPACE_SIGA, f"factura_julio_{numero_medidor}"))
        facturas_list.append({
            'id': factura_id,
            'numero_factura': f"FAC-JUL-{numero_medidor}",
            'id_socio': socio['id'],
            'id_medidor': medidor_id,
            'id_periodo': PERIODO_JULIO_ID,
            'id_lectura': lectura_id,
            'es_tercera_edad': socio['es_tercera_edad'],
            'valor_base': 0.0,
            'consumo_m3': 0.0,
            'excedente_m3': 0.0,
            'valor_excedente': 0.0,
            'valor_alcantarillado': 0.0,
            'valor_multas': 0.0,
            'valor_deuda_anterior': deuda_pen,
            'total_mes': deuda_pen,
            'total_pagar': deuda_pen,
            'monto_pagado': 0.0,
            'saldo_pendiente': deuda_pen,
            'estado_pago': 'PENDIENTE',
            'fecha_vencimiento': '2026-08-15',
            'version': 1
        })
        
    # 3. Multas y Rubros al corte de Julio
    if multas_val > 0:
        multa_id = str(uuid.uuid5(NAMESPACE_SIGA, f"multa_julio_{row_idx}_{socio['id']}"))
        tipo_rubro = 'ASAMBLEA' if multas_val <= 10 else 'MINGA'
        multas_list.append({
            'id': multa_id,
            'id_socio': socio['id'],
            'id_periodo': PERIODO_JULIO_ID,
            'tipo_rubro': tipo_rubro,
            'monto': multas_val,
            'motivo': f"Multas acumuladas al corte de Julio 2026 ({alias if alias else numero_medidor})",
            'pagado': False,
            'id_factura': None,
            'created_at': '2026-07-31T23:59:59.000Z'
        })

result = {
    'version': '2.0.0-corte-julio',
    'timestamp': datetime.datetime.now().isoformat(),
    'sectores': list(sectores_map.values()),
    'periodos': periodos,
    'socios': list(socios_dict.values()),
    'medidores': medidores_list,
    'lecturas': lecturas_list,
    'facturas': facturas_list,
    'multas_rubros': multas_list
}

output_path = 'scripts/seed_normalizado.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(result, f, indent=2, ensure_ascii=False)

print("\n" + "="*60)
print(" RESUMEN DE EXTRACCIÓN Y NORMALIZACIÓN DE SEMILLA ")
print("="*60)
print(f" Sectores oficiales:         {len(result['sectores'])}")
print(f" Períodos (Julio/Agosto):    {len(result['periodos'])}")
print(f" Socios únicos (Titulares):  {len(result['socios'])} (SOC-0001 a SOC-{str(len(result['socios'])).zfill(4)})")
print(f" Medidores (Acometidas):     {len(result['medidores'])}")
print(f" Lecturas de corte (Julio):  {len(result['lecturas'])}")
print(f" Facturas en mora (Julio):   {len(result['facturas'])}")
print(f" Multas en mora (Julio):     {len(result['multas_rubros'])}")
print(f" Total deuda acumulada:      ${sum(f['total_pagar'] for f in result['facturas']):,.2f}")
print(f" Total multas acumuladas:    ${sum(m['monto'] for m in result['multas_rubros']):,.2f}")
print("="*60)
print(f" Archivo generado: {output_path}")
